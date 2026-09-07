# -*- coding: utf-8 -*-
"""Core training, saving, loading, and imputation utilities for the groundwater paper attachment."""

import json
import os
from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd


@dataclass
class ColumnInfo:
    date_col: str
    well_cols: List[str]
    external_cols: List[str]


def clean_column_names(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c).strip() for c in out.columns]
    return out


def identify_columns(df: pd.DataFrame, external_cols: Optional[Sequence[str]] = None) -> ColumnInfo:
    """Identify date, well and external-variable columns.

    Rules
    -----
    - Date column: name equals ``date`` ignoring case.
    - Well columns: column name contains ``well`` ignoring case.
    - External columns: either user-specified, or all remaining non-date/non-well columns.
    """
    columns = [str(c).strip() for c in df.columns]
    date_candidates = [c for c in columns if c.lower() == "date"]
    if not date_candidates:
        raise ValueError("The input table must contain a date column named 'date' (case-insensitive).")
    date_col = date_candidates[0]

    well_cols = [c for c in columns if "well" in c.lower()]
    if not well_cols:
        raise ValueError("No well columns were detected. Well column names must contain 'well'.")

    if external_cols is None:
        ext_cols = [c for c in columns if c != date_col and c not in well_cols]
    else:
        ext_cols = [str(c).strip() for c in external_cols if str(c).strip()]
        missing = [c for c in ext_cols if c not in columns]
        if missing:
            raise ValueError(f"The following --external-cols are not in the table: {missing}")

    return ColumnInfo(date_col=date_col, well_cols=well_cols, external_cols=ext_cols)


def ensure_datetime_sorted(df: pd.DataFrame, date_col: str) -> pd.DataFrame:
    out = clean_column_names(df)
    out[date_col] = pd.to_datetime(out[date_col], errors="coerce")
    out = out[out[date_col].notna()].copy()
    out = out.sort_values(date_col).reset_index(drop=True)
    return out


def infer_time_delta(dates: pd.Series) -> Optional[pd.Timedelta]:
    dt = pd.to_datetime(dates, errors="coerce").dropna().sort_values().diff().dropna()
    dt = dt[dt > pd.Timedelta(0)]
    if dt.empty:
        return None
    mode_vals = dt.mode()
    return mode_vals.iloc[0] if len(mode_vals) else dt.median()


def is_effectively_not_missing(s: pd.Series) -> pd.Series:
    """Treat NaN and blank strings as missing values."""
    return s.notna() & (~s.astype(str).str.strip().eq(""))


def build_complete_mask(df: pd.DataFrame, required_cols: Sequence[str]) -> pd.Series:
    if not required_cols:
        raise ValueError("required_cols cannot be empty.")
    mask = pd.Series(True, index=df.index)
    for col in required_cols:
        if col not in df.columns:
            raise ValueError(f"Required column not found: {col}")
        mask &= is_effectively_not_missing(df[col])
    return mask


def boolean_true_segments(mask: pd.Series) -> pd.DataFrame:
    """Return continuous True segments in a boolean mask."""
    arr = mask.to_numpy()
    rows = []
    start = None
    for i, flag in enumerate(arr):
        if flag and start is None:
            start = i
        if (not flag or i == len(arr) - 1) and start is not None:
            end = i if flag and i == len(arr) - 1 else i - 1
            rows.append({"start_pos": start, "end_pos": end, "length": end - start + 1})
            start = None
    return pd.DataFrame(rows, columns=["start_pos", "end_pos", "length"])


def select_longest_complete_block(
    df: pd.DataFrame,
    date_col: str,
    required_cols: Sequence[str],
    sort_by_date: bool = True,
) -> Tuple[int, int, pd.DataFrame, pd.DataFrame]:
    """Select the longest block where all required columns are non-missing.

    Time-interval continuity is intentionally not enforced. This keeps the
    selection rule simple and consistent with the training data requirement:
    the selected rows must be complete for all modelling variables.
    """
    work = clean_column_names(df)
    work[date_col] = pd.to_datetime(work[date_col], errors="coerce")
    work = work[work[date_col].notna()].copy()
    if sort_by_date:
        work = work.sort_values(date_col).reset_index(drop=True)
    else:
        work = work.reset_index(drop=True)

    complete_mask = build_complete_mask(work, required_cols)
    segments = boolean_true_segments(complete_mask)
    if segments.empty:
        raise ValueError("No complete block was found for the required modelling columns.")

    segments = segments.sort_values(["length", "start_pos"], ascending=[False, True]).reset_index(drop=True)
    start = int(segments.loc[0, "start_pos"])
    end = int(segments.loc[0, "end_pos"])
    block = work.iloc[start : end + 1].reset_index(drop=True)
    return start, end, block, segments


def read_json(path: str) -> Dict:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def write_json(obj: Dict, path: str) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)


def safe_float(value, default: float = 0.0) -> float:
    try:
        if value is None or pd.isna(value):
            return default
        return float(value)
    except Exception:
        return default


def contiguous_nan_segments(series: pd.Series) -> List[Dict[str, int]]:
    """Return contiguous NaN segments as inclusive integer row ranges."""
    is_nan = series.isna().to_numpy()
    segments: List[Dict[str, int]] = []
    start: Optional[int] = None
    for i, flag in enumerate(is_nan):
        if flag and start is None:
            start = i
        if (not flag or i == len(is_nan) - 1) and start is not None:
            end = i if flag and i == len(is_nan) - 1 else i - 1
            segments.append({"start": start, "end": end, "length": end - start + 1})
            start = None
    return segments


# =============================================================================
# Three-module model training code
# =============================================================================

import os
import glob
import random
import warnings
import argparse
import sys
import pickle
import json
from itertools import combinations
import copy
import numpy as np
import pandas as pd
import matplotlib

matplotlib.use('Agg')

warnings.filterwarnings('ignore')

SEED = 42
os.environ['PYTHONHASHSEED'] = str(SEED)
os.environ['TF_CPP_MIN_LOG_LEVEL'] = '2'
os.environ['TF_DETERMINISTIC_OPS'] = '1'

import tensorflow as tf

try:
    tf.config.experimental.enable_op_determinism()
except AttributeError:
    pass

gpus = tf.config.experimental.list_physical_devices('GPU')
if gpus:
    try:
        for gpu in gpus:
            tf.config.experimental.set_memory_growth(gpu, True)
    except RuntimeError as e:
        print(f"TensorFlow GPU memory configuration error: {e}")

from tensorflow.keras.models import Sequential, load_model
from tensorflow.keras.layers import LSTM as TF_LSTM, Dropout, Dense, Bidirectional
from tensorflow.keras.callbacks import EarlyStopping
from sklearn.preprocessing import StandardScaler, MinMaxScaler

import torch
import torch.nn as nn
from torch.utils.data import DataLoader, TensorDataset
import torch.optim as optim

from scipy.optimize import minimize
from sklearn.model_selection import train_test_split
import xgboost as xgb

random.seed(SEED)
np.random.seed(SEED)
tf.random.set_seed(SEED)

torch.manual_seed(SEED)
if torch.cuda.is_available():
    torch.cuda.manual_seed(SEED)
    torch.cuda.manual_seed_all(SEED)
torch.backends.cudnn.deterministic = True
torch.backends.cudnn.benchmark = False

GLOBAL_DATA_FILE = 'data.xlsx'

CONTEXT_LENGTH = 8
TS_FORECAST_STEPS = [1, 2, 3, 6, 12, 24, 48, 72, 120, 168, 336, 504, 720]
TS_PARAM_GRID = {
    'lstm_units': [16, 32],
    'dropout_rate': [0.1, 0.2],
    'learning_rate': [0.001, 0.0005],
    'batch_size': [32]
}

TRAIN_END_IDX = 2000
VAL_END_IDX = 3000
MISSING_PROB = 0.2

EXT_HYPERPARAMS = {
    'n_estimators': [50, 100, 200],
    'max_depth': [3, 5, 7],
    'learning_rate': [0.01, 0.05, 0.1]
}
DEVICE = torch.device("cuda" if torch.cuda.is_available() else "cpu")


def get_model_save_dir(base_out_dir):
    model_dir = os.path.join(base_out_dir, "saved_models")
    os.makedirs(model_dir, exist_ok=True)
    return model_dir


def compute_metrics_np(y_true, y_pred):
    if len(y_true) == 0:
        return np.nan, np.nan, np.nan, np.nan
    rmse = np.sqrt(np.mean((y_true - y_pred) ** 2))
    mean_true = np.mean(y_true)
    denom = np.sum((y_true - mean_true) ** 2)
    nse = 1 - np.sum((y_true - y_pred) ** 2) / denom if denom != 0 else -np.inf

    mae = np.mean(np.abs(y_true - y_pred))
    # Avoid division by zero in MAPE.
    non_zero_mask = y_true != 0
    if np.sum(non_zero_mask) > 0:
        mape = np.mean(np.abs((y_true[non_zero_mask] - y_pred[non_zero_mask]) / y_true[non_zero_mask])) * 100
    else:
        mape = np.nan

    return rmse, nse, mae, mape


# ==============================================================================
# Module 1: univariate target-well Bi-LSTM interpolation.
# ==============================================================================
def step1_univariate_sequence(target_well, lstm_results_dir, base_out_dir):
    print("\n" + "=" * 50)
    print(f">>> Module 1: univariate Bi-LSTM interpolation - target well: {target_well}")
    print("=" * 50)

    def load_data(file_path):
        df = pd.read_excel(file_path)
        date_col = [col for col in df.columns if col.lower() == 'date'][0]
        df[date_col] = pd.to_datetime(df[date_col])
        df = df.sort_values(date_col).reset_index(drop=True)

        # Fix: use the raw unified index to define absolute split dates.
        train_end_date = df.loc[TRAIN_END_IDX, date_col]
        val_end_date = df.loc[VAL_END_IDX, date_col]

        df = df.set_index(date_col)
        series = df[target_well].astype(float)
        dates = series.index
        values = series.values

        # The scaler is fitted only on the training period defined by absolute dates.
        train_mask = dates < train_end_date
        scaler = StandardScaler().fit(values[train_mask].reshape(-1, 1))
        scaled = scaler.transform(values.reshape(-1, 1)).flatten()
        return scaled, dates, scaler, train_end_date, val_end_date

    def create_samples(seq, missing_len, stride=1):
        X, y, start_indices = [], [], []
        window_size = CONTEXT_LENGTH * 2 + missing_len
        for i in range(0, len(seq) - window_size + 1, stride):
            past = seq[i: i + CONTEXT_LENGTH]
            future = seq[i + CONTEXT_LENGTH + missing_len: i + window_size]
            target = seq[i + CONTEXT_LENGTH: i + CONTEXT_LENGTH + missing_len]
            X.append(np.concatenate([past, future]))
            y.append(target)
            start_indices.append(i)
        if not X:
            return np.array([]), np.array([]), np.array([])
        return np.array(X).reshape(-1, CONTEXT_LENGTH * 2, 1), np.array(y).reshape(-1, missing_len), np.array(
            start_indices)

    def build_bilstm_model(units, dropout, lr, out_steps):
        model = Sequential([
            Bidirectional(TF_LSTM(units, return_sequences=True), input_shape=(CONTEXT_LENGTH * 2, 1)),
            Dropout(dropout),
            Bidirectional(TF_LSTM(units, return_sequences=False)),
            Dropout(dropout),
            Dense(out_steps)
        ])
        model.compile(optimizer=tf.keras.optimizers.Adam(lr), loss='mse')
        return model

    def process_predictions(X, y_true, window_indices, model, base_offset, d, dataset_label, scaler, dates):
        if len(X) == 0:
            return pd.DataFrame()
        preds = model.predict(X, verbose=0)

        preds_orig = scaler.inverse_transform(preds.reshape(-1, 1)).reshape(-1, d)
        targets_orig = scaler.inverse_transform(y_true.reshape(-1, 1)).reshape(-1, d)

        records = []
        for k in range(len(preds)):
            window_start = window_indices[k]
            for j in range(d):
                abs_idx = base_offset + window_start + CONTEXT_LENGTH + j
                records.append({
                    'step': d,
                    'idx': abs_idx,
                    'date': dates[abs_idx],
                    'true': targets_orig[k, j],
                    'pred': preds_orig[k, j],
                    'Dataset': dataset_label,
                    'abs_pos': j + 1
                })
        return pd.DataFrame(records)

    scaled, dates, scaler, train_end_date, val_end_date = load_data(GLOBAL_DATA_FILE)

    # Fix: split by date masks and compute offsets to align with dates.
    train_mask = dates < train_end_date
    val_mask = (dates >= train_end_date) & (dates < val_end_date)
    test_mask = dates >= val_end_date

    train_data = scaled[train_mask]
    val_data = scaled[val_mask]
    test_data = scaled[test_mask]

    train_start_offset = 0
    val_start_offset = int(train_mask.sum())
    test_start_offset = int(train_mask.sum() + val_mask.sum())

    model_save_dir = get_model_save_dir(base_out_dir)
    univ_dir = os.path.join(model_save_dir, "univariate")
    os.makedirs(univ_dir, exist_ok=True)
    with open(os.path.join(univ_dir, f"scaler_{target_well}.pkl"), 'wb') as f:
        pickle.dump(scaler, f)

    d_search = TS_FORECAST_STEPS[0]
    X_tr_s, y_tr_s, _ = create_samples(train_data, d_search, stride=1)
    X_va_s, y_va_s, _ = create_samples(val_data, d_search, stride=d_search)

    best_params = None
    best_val_rmse = np.inf

    if len(X_tr_s) > 0 and len(X_va_s) > 0:
        for units in TS_PARAM_GRID['lstm_units']:
            for do in TS_PARAM_GRID['dropout_rate']:
                for lr in TS_PARAM_GRID['learning_rate']:
                    for bs in TS_PARAM_GRID['batch_size']:
                        model = build_bilstm_model(units, do, lr, d_search)
                        es = EarlyStopping(monitor='val_loss', patience=10, restore_best_weights=True, verbose=0)
                        model.fit(X_tr_s, y_tr_s, validation_data=(X_va_s, y_va_s), epochs=100, batch_size=bs,
                                  callbacks=[es], verbose=0)
                        preds_va = model.predict(X_va_s, verbose=0)
                        val_rmse = np.sqrt(np.mean((y_va_s - preds_va) ** 2))
                        if val_rmse < best_val_rmse:
                            best_val_rmse = val_rmse
                            best_params = {'lstm_units': units, 'dropout_rate': do, 'learning_rate': lr,
                                           'batch_size': bs}
        print(f"Using d={d_search} selected global hyperparameters: {best_params}, Val RMSE = {best_val_rmse:.6f}")
    else:
        best_params = {'lstm_units': 32, 'dropout_rate': 0.1, 'learning_rate': 0.001, 'batch_size': 32}
        print("Warning: not enough data for hyperparameter search; default parameters are used.")

    test_metrics = []

    for d in TS_FORECAST_STEPS:
        X_tr, y_tr, ind_tr = create_samples(train_data, d, stride=1)
        X_va, y_va, ind_va = create_samples(val_data, d, stride=d)
        X_te, y_te, ind_te = create_samples(test_data, d, stride=d)

        if len(X_tr) == 0:
            print(f"Step {d} is too large for valid training windows; skipped.")
            continue

        model = build_bilstm_model(best_params['lstm_units'], best_params['dropout_rate'], best_params['learning_rate'],
                                   d)
        es = EarlyStopping(monitor='val_loss', patience=10, restore_best_weights=True, verbose=0)

        if len(X_va) > 0:
            model.fit(X_tr, y_tr, validation_data=(X_va, y_va), epochs=100, batch_size=best_params['batch_size'],
                      callbacks=[es], verbose=0)
        else:
            model.fit(X_tr, y_tr, epochs=100, batch_size=best_params['batch_size'], verbose=0)

        model.save(os.path.join(univ_dir, f"model_{target_well}_step_{d}.keras"))

        df_tr = process_predictions(X_tr, y_tr, ind_tr, model, train_start_offset, d, 'train', scaler, dates)
        df_va = process_predictions(X_va, y_va, ind_va, model, val_start_offset, d, 'val', scaler, dates)
        df_te = process_predictions(X_te, y_te, ind_te, model, test_start_offset, d, 'test', scaler, dates)

        df_step = pd.concat([df_tr, df_va, df_te], ignore_index=True)
        if not df_step.empty:
            df_step.to_csv(os.path.join(lstm_results_dir, f'predictions_step_{d:03d}.csv'), index=False)

            mask_test = df_step['Dataset'] == 'test'
            if mask_test.sum() > 0:
                rmse, nse, mae, mape = compute_metrics_np(df_step.loc[mask_test, 'true'].values,
                                                          df_step.loc[mask_test, 'pred'].values)
                test_metrics.append({'Step': d, 'RMSE': rmse, 'NSE': nse, 'MAE': mae, 'MAPE': mape})
                print(f"d={d:03d} completed -> Test RMSE: {rmse:.4f}, NSE: {nse:.4f}, MAE: {mae:.4f}, MAPE: {mape:.4f}%")

    pd.DataFrame(test_metrics).to_csv(os.path.join(base_out_dir, 'step_test_metrics.csv'), index=False)
    print(f"Module 1 finished. Results saved to {lstm_results_dir}")


# ==============================================================================
# Module 2: XGBoost instantaneous external-feature prediction.
# ==============================================================================
def step2_multivariate_external(target_well, feature_cols, lstm_results_dir, base_out_dir):
    print("\n" + "=" * 50)
    print(f">>> Module 2: external-feature selection and XGBoost training")
    print("=" * 50)

    if target_well in feature_cols:
        raise ValueError(f"Critical error: target well {target_well} was included as an external predictor, which is invalid.")

    df_raw = pd.read_excel(GLOBAL_DATA_FILE)
    date_col = [col for col in df_raw.columns if col.lower() == 'date'][0]
    df_raw[date_col] = pd.to_datetime(df_raw[date_col])
    df_raw = df_raw.sort_values(date_col).reset_index(drop=True)

    # Fix: use absolute split dates aligned with Module 1.
    train_end_date = df_raw.loc[TRAIN_END_IDX, date_col]
    val_end_date = df_raw.loc[VAL_END_IDX, date_col]
    dates_all = df_raw[date_col].values

    dates_dt = df_raw[date_col]
    train_mask = dates_dt < train_end_date
    val_mask = (dates_dt >= train_end_date) & (dates_dt < val_end_date)

    df_imputed = df_raw.copy()
    augmented_features = []
    # These training means are reused during deployment when features are missing.
    m2_feature_means = {}

    np.random.seed(SEED)
    for col in feature_cols:
        mask = np.random.binomial(1, 1 - MISSING_PROB, size=len(df_imputed))
        mask_col_name = f"{col}_mask"
        df_imputed[mask_col_name] = mask

        # Fix: compute means and fit scalers strictly on the training mask.
        train_vals = df_imputed.loc[train_mask, col]
        train_mean = train_vals[mask[train_mask] == 1].mean()
        if np.isnan(train_mean): train_mean = 0
        m2_feature_means[col] = float(train_mean)

        df_imputed[col] = np.where(mask == 1, df_imputed[col], train_mean)
        augmented_features.append(col)
        augmented_features.append(mask_col_name)

    all_feature_combinations = []
    for r in range(1, len(feature_cols) + 1):
        all_feature_combinations.extend(list(combinations(feature_cols, r)))

    model_save_dir = get_model_save_dir(base_out_dir)
    multiv_dir = os.path.join(model_save_dir, "multivariate")
    os.makedirs(multiv_dir, exist_ok=True)

    best_global_val_loss = float('inf')
    best_global_features = None
    best_global_model = None
    best_global_scaler_X = None
    best_global_scaler_y = None

    idx_seq = np.arange(len(df_imputed))

    print(f"Generated {len(all_feature_combinations)} base feature combinations. Starting search...")

    for comb_idx, base_feat_comb in enumerate(all_feature_combinations):
        current_features = []
        for f in base_feat_comb:
            current_features.extend([f, f"{f}_mask"])

        X_all_raw = df_imputed[current_features].values
        y_all_raw = df_imputed[[target_well]].values

        scaler_X = MinMaxScaler().fit(X_all_raw[train_mask])
        scaler_y = MinMaxScaler().fit(y_all_raw[train_mask])

        X_all_scaled = scaler_X.transform(X_all_raw)
        y_all_scaled = scaler_y.transform(y_all_raw).flatten()

        X_train, y_train = X_all_scaled[train_mask], y_all_scaled[train_mask]
        X_val, y_val = X_all_scaled[val_mask], y_all_scaled[val_mask]

        for n_est in EXT_HYPERPARAMS['n_estimators']:
            for md in EXT_HYPERPARAMS['max_depth']:
                for lr in EXT_HYPERPARAMS['learning_rate']:
                    model = xgb.XGBRegressor(n_estimators=n_est, max_depth=md, learning_rate=lr, random_state=SEED,
                                             n_jobs=-1)
                    model.fit(X_train, y_train)
                    val_preds = model.predict(X_val)
                    val_rmse = np.sqrt(np.mean((y_val - val_preds) ** 2))

                    if val_rmse < best_global_val_loss:
                        best_global_val_loss = val_rmse
                        best_global_model = model
                        best_global_features = current_features
                        best_global_scaler_X = scaler_X
                        best_global_scaler_y = scaler_y

    print(f"Best feature combination selected: {best_global_features}")

    X_best_all_raw = df_imputed[best_global_features].values
    X_best_scaled = best_global_scaler_X.transform(X_best_all_raw)
    preds_scaled = best_global_model.predict(X_best_scaled).reshape(-1, 1)

    preds_real = best_global_scaler_y.inverse_transform(preds_scaled)
    y_true_real = df_imputed[[target_well]].values

    safe_feat_str = "Best_XGB_Comb"
    best_global_model.save_model(os.path.join(multiv_dir, f"model_{safe_feat_str}.json"))
    with open(os.path.join(multiv_dir, f"scaler_X_{safe_feat_str}.pkl"), 'wb') as f:
        pickle.dump(best_global_scaler_X, f)
    with open(os.path.join(multiv_dir, f"scaler_y_{safe_feat_str}.pkl"), 'wb') as f:
        pickle.dump(best_global_scaler_y, f)

    # Save M2 deployment metadata: feature order, means, model path, and scaler paths.
    m2_metadata = {
        "target_well": target_well,
        "feature_cols_candidate": list(feature_cols),
        "features_used": list(best_global_features),
        "feature_means": m2_feature_means,
        "model_path": os.path.abspath(os.path.join(multiv_dir, f"model_{safe_feat_str}.json")),
        "scaler_X_path": os.path.abspath(os.path.join(multiv_dir, f"scaler_X_{safe_feat_str}.pkl")),
        "scaler_y_path": os.path.abspath(os.path.join(multiv_dir, f"scaler_y_{safe_feat_str}.pkl")),
        "train_end_date": str(train_end_date),
        "val_end_date": str(val_end_date)
    }
    with open(os.path.join(multiv_dir, "m2_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(m2_metadata, f, ensure_ascii=False, indent=2)

    # Fix: assign dataset labels by boolean masks.
    dataset_labels = np.where(train_mask, 'train', np.where(val_mask, 'val', 'test'))

    results_df = pd.DataFrame({
        'Index': idx_seq, 'Date': dates_all, 'Dataset': dataset_labels,
        'True_Value': y_true_real.flatten(), 'Pred_Value': preds_real.flatten(),
        'Features_Used': "|".join(best_global_features)
    })

    feat_df = pd.DataFrame(X_best_all_raw, columns=best_global_features)
    results_df = pd.concat([results_df, feat_df], axis=1)
    results_df = results_df.iloc[CONTEXT_LENGTH:].reset_index(drop=True)

    best_file_name = os.path.join(lstm_results_dir, f"Best_Comb_M2.csv")
    results_df.to_csv(best_file_name, index=False, encoding='utf-8-sig')


# ==============================================================================
# Module 3: gating network fusion.
# ==============================================================================
class FusionGatingNet(nn.Module):
    def __init__(self, input_dim):
        super(FusionGatingNet, self).__init__()
        self.net = nn.Sequential(
            nn.Linear(input_dim, 64),
            nn.ReLU(),
            nn.Dropout(0.2),
            nn.Linear(64, 32),
            nn.ReLU(),
            nn.Linear(32, 1),
            nn.Sigmoid()
        )

    def forward(self, x):
        return self.net(x).squeeze(-1)


def step3_fusion(lstm_results_dir, base_out_dir, fuse_output_dir_details, fuse_summary_filepath):
    print("\n" + "=" * 50)
    print(">>> Module 3: neural gating fusion with strict index alignment")
    print("=" * 50)

    m1_files = sorted(glob.glob(os.path.join(lstm_results_dir, "predictions_step_*.csv")))
    m2_file = os.path.join(lstm_results_dir, "Best_Comb_M2.csv")

    if not m1_files or not os.path.exists(m2_file):
        print("[Error] Required prediction files were not found. Fusion stopped.")
        return

    # Read directly; avoid forced Date parsing issues.
    df2 = pd.read_csv(m2_file)
    features_used = df2['Features_Used'].iloc[0].split('|')

    all_merged_data = []
    for m1_file in m1_files:
        df1 = pd.read_csv(m1_file)

        # Align Module 1 and Module 2 predictions by physical index.
        merged = pd.merge(df2, df1, left_on='Index', right_on='idx', how='inner', suffixes=('', '_m1'))
        if not merged.empty:
            all_merged_data.append(merged)

    if not all_merged_data:
        print("[Error] No valid aligned samples after merging. Fusion stopped.")
        return

    full_df = pd.concat(all_merged_data, ignore_index=True)

    # Sort by index and step to preserve temporal order.
    full_df.sort_values(by=['Index', 'step'], inplace=True)
    full_df.reset_index(drop=True, inplace=True)

    target_dataset_col = 'Dataset' if 'Dataset' in full_df.columns else 'Dataset_m1'

    step_values = full_df['step'].values.astype(float).reshape(-1, 1)
    abs_pos_values = full_df['abs_pos'].values.astype(float).reshape(-1, 1)
    env_features = full_df[features_used].values
    X_raw = np.hstack((step_values, abs_pos_values, env_features))

    is_val_all = full_df[target_dataset_col].str.lower() == 'val'
    is_test = full_df[target_dataset_col].str.lower() == 'test'
    is_train = full_df[target_dataset_col].str.lower() == 'train'

    if is_val_all.sum() == 0:
        print("[Error] Validation set is empty; the gating network cannot be trained.")
        return

    val_df = full_df[is_val_all].copy()
    X_val_raw = X_raw[is_val_all]
    y_val_true = full_df['True_Value'].values[is_val_all].astype(np.float32)
    pred_m1_val = full_df['pred'].values[is_val_all].astype(np.float32)
    pred_m2_val = full_df['Pred_Value'].values[is_val_all].astype(np.float32)

    # Train/validate the gating network using a temporal split within the validation period.
    gate_train_size = int(len(val_df) * 0.8)
    gate_train_idx = np.arange(gate_train_size)
    gate_val_idx = np.arange(gate_train_size, len(val_df))

    X_gate_train_raw = X_val_raw[gate_train_idx]
    X_gate_val_raw = X_val_raw[gate_val_idx]
    y_gate_train = y_val_true[gate_train_idx]
    y_gate_val = y_val_true[gate_val_idx]
    m1_gate_train = pred_m1_val[gate_train_idx]
    m1_gate_val = pred_m1_val[gate_val_idx]
    m2_gate_train = pred_m2_val[gate_train_idx]
    m2_gate_val = pred_m2_val[gate_val_idx]

    scaler_gate = MinMaxScaler()
    scaler_gate.fit(X_gate_train_raw)

    X_gate_train_scaled = scaler_gate.transform(X_gate_train_raw)
    X_gate_val_scaled = scaler_gate.transform(X_gate_val_raw)

    X_test_raw = X_raw[is_test]
    X_test_scaled = scaler_gate.transform(X_test_raw) if is_test.sum() > 0 else np.array([])
    pred_m1_test = full_df['pred'].values[is_test].astype(np.float32)
    pred_m2_test = full_df['Pred_Value'].values[is_test].astype(np.float32)

    X_train_raw = X_raw[is_train]
    X_train_scaled = scaler_gate.transform(X_train_raw)
    pred_m1_train = full_df['pred'].values[is_train].astype(np.float32)
    pred_m2_train = full_df['Pred_Value'].values[is_train].astype(np.float32)

    X_train_t = torch.tensor(X_gate_train_scaled, dtype=torch.float32).to(DEVICE)
    y_true_train_t = torch.tensor(y_gate_train, dtype=torch.float32).to(DEVICE)
    m1_train_t = torch.tensor(m1_gate_train, dtype=torch.float32).to(DEVICE)
    m2_train_t = torch.tensor(m2_gate_train, dtype=torch.float32).to(DEVICE)

    X_val_t = torch.tensor(X_gate_val_scaled, dtype=torch.float32).to(DEVICE)
    y_true_val_t = torch.tensor(y_gate_val, dtype=torch.float32).to(DEVICE)
    m1_val_t = torch.tensor(m1_gate_val, dtype=torch.float32).to(DEVICE)
    m2_val_t = torch.tensor(m2_gate_val, dtype=torch.float32).to(DEVICE)

    train_dataset = TensorDataset(X_train_t, y_true_train_t, m1_train_t, m2_train_t)
    train_loader = DataLoader(train_dataset, batch_size=256, shuffle=True)

    input_dim = X_train_t.shape[1]
    gating_net = FusionGatingNet(input_dim).to(DEVICE)
    optimizer = optim.Adam(gating_net.parameters(), lr=0.001)

    best_val_loss = float('inf')
    best_model_state = None
    epochs = 150
    print(f"Starting gating-network training. Gate_Train samples: {X_train_t.shape[0]}; Gate_Val samples: {X_val_t.shape[0]}")

    for epoch in range(epochs):
        gating_net.train()
        for batch_X, batch_y, batch_m1, batch_m2 in train_loader:
            optimizer.zero_grad()
            w_batch = gating_net(batch_X)
            fused_batch = w_batch * batch_m1 + (1.0 - w_batch) * batch_m2
            loss = torch.mean((batch_y - fused_batch) ** 2)
            loss.backward()
            optimizer.step()

        gating_net.eval()
        with torch.no_grad():
            w_val = gating_net(X_val_t)
            fused_val = w_val * m1_val_t + (1.0 - w_val) * m2_val_t
            val_loss = torch.mean((y_true_val_t - fused_val) ** 2).item()
            if val_loss < best_val_loss:
                best_val_loss = val_loss
                best_model_state = copy.deepcopy(gating_net.state_dict())

    gating_net.load_state_dict(best_model_state)

    gating_net.eval()
    with torch.no_grad():
        train_t = torch.tensor(X_train_scaled, dtype=torch.float32).to(DEVICE)
        w_train = gating_net(train_t).cpu().numpy().flatten()
        fused_train = w_train * pred_m1_train.flatten() + (1.0 - w_train) * pred_m2_train.flatten()

        val_all_scaled = scaler_gate.transform(X_val_raw)
        val_all_t = torch.tensor(val_all_scaled, dtype=torch.float32).to(DEVICE)
        w_val_all = gating_net(val_all_t).cpu().numpy().flatten()
        fused_val_all = w_val_all * pred_m1_val.flatten() + (1.0 - w_val_all) * pred_m2_val.flatten()

        if is_test.sum() > 0:
            test_t = torch.tensor(X_test_scaled, dtype=torch.float32).to(DEVICE)
            w_test = gating_net(test_t).cpu().numpy().flatten()
            fused_test = w_test * pred_m1_test.flatten() + (1.0 - w_test) * pred_m2_test.flatten()
        else:
            w_test, fused_test = np.array([]), np.array([])

    full_df['Fused_Pred'] = np.nan
    full_df.loc[is_train, 'Fused_Pred'] = fused_train.flatten()
    full_df.loc[is_val_all, 'Fused_Pred'] = fused_val_all.flatten()
    if is_test.sum() > 0:
        full_df.loc[is_test, 'Fused_Pred'] = fused_test.flatten()

    full_df['Model_1_Weight'] = np.nan
    full_df.loc[is_train, 'Model_1_Weight'] = w_train.flatten()
    full_df.loc[is_val_all, 'Model_1_Weight'] = w_val_all.flatten()
    if is_test.sum() > 0:
        full_df.loc[is_test, 'Model_1_Weight'] = w_test.flatten()
    full_df['Model_2_Weight'] = 1.0 - full_df['Model_1_Weight']

    # =========================================================
    # Export fusion details:
    # Gate input remains: step + abs_pos + M2-selected features.
    # Only additional output fields are added; the training logic is unchanged.
    # =========================================================

    out_cols = [
        'Index', 'Date', 'step', 'abs_pos', target_dataset_col,
        'True_Value',
        'pred', 'Pred_Value',
        'Model_1_Weight', 'Model_2_Weight', 'Fused_Pred'
    ]

    # Output the M2-selected features actually used by the gate.
    gate_feature_cols = [c for c in features_used if c in full_df.columns]

    out_df_all = full_df[out_cols + gate_feature_cols].copy()

    out_df_all.rename(columns={
        'pred': 'Model_1_Pred_LSTM',
        'Pred_Value': 'Model_2_Pred_XGB',
        target_dataset_col: 'Dataset'
    }, inplace=True)

    # Missing-segment metadata.
    out_df_all['segment_start'] = out_df_all['Index'] - out_df_all['abs_pos'] + 1
    out_df_all['segment_id'] = (
            out_df_all['step'].astype(str)
            + '_'
            + out_df_all['segment_start'].astype(str)
    )

    # Relative position within the missing segment.
    out_df_all['pos_ratio'] = out_df_all['abs_pos'] / out_df_all['step']

    # =========================================================
    # Add point-wise errors and fusion gains.
    # =========================================================

    for model_name, pred_col in [
        ('M1', 'Model_1_Pred_LSTM'),
        ('M2', 'Model_2_Pred_XGB'),
        ('Fused', 'Fused_Pred')
    ]:
        out_df_all[f'Err_{model_name}'] = out_df_all[pred_col] - out_df_all['True_Value']
        out_df_all[f'AE_{model_name}'] = np.abs(out_df_all[f'Err_{model_name}'])
        out_df_all[f'SE_{model_name}'] = out_df_all[f'Err_{model_name}'] ** 2

    # Positive values indicate that fusion is better than the corresponding single model.
    out_df_all['Gain_AE_vs_M1'] = out_df_all['AE_M1'] - out_df_all['AE_Fused']
    out_df_all['Gain_AE_vs_M2'] = out_df_all['AE_M2'] - out_df_all['AE_Fused']

    out_df_all['Gain_SE_vs_M1'] = out_df_all['SE_M1'] - out_df_all['SE_Fused']
    out_df_all['Gain_SE_vs_M2'] = out_df_all['SE_M2'] - out_df_all['SE_Fused']

    # Which single model is better at the current point.
    out_df_all['Better_Model'] = np.where(
        out_df_all['AE_M1'] <= out_df_all['AE_M2'],
        'M1',
        'M2'
    )

    # Which model the gate weights favor.
    out_df_all['Gate_Choose'] = np.where(
        out_df_all['Model_1_Weight'] >= 0.5,
        'M1',
        'M2'
    )

    # Whether the gate favors the point-wise better single model.
    out_df_all['Gate_Routing_Correct'] = (
            out_df_all['Better_Model'] == out_df_all['Gate_Choose']
    ).astype(int)

    # =========================================================
    # Save all-set and test-only fusion details.
    # =========================================================

    out_df_all.to_csv(
        os.path.join(fuse_output_dir_details, "Gating_Fusion_All_With_GateInputs.csv"),
        index=False,
        encoding='utf-8-sig'
    )

    out_df_test_only = out_df_all[out_df_all['Dataset'].str.lower() == 'test'].copy()

    out_df_test_only.to_csv(
        os.path.join(fuse_output_dir_details, "Gating_Fusion_Test_With_GateInputs.csv"),
        index=False,
        encoding='utf-8-sig'
    )

    # =========================================================
    # Test-set missing-segment-level summary.
    # One row represents one missing segment rather than one point.
    # =========================================================

    segment_summary_rows = []

    for seg_id, df_seg in out_df_test_only.groupby('segment_id'):
        row = {
            'segment_id': seg_id,
            'step': df_seg['step'].iloc[0],
            'segment_start': df_seg['segment_start'].iloc[0],
            'start_date': df_seg['Date'].iloc[0],
            'end_date': df_seg['Date'].iloc[-1],
            'n_points': len(df_seg),
            'mean_Model_1_Weight': df_seg['Model_1_Weight'].mean(),
            'std_Model_1_Weight': df_seg['Model_1_Weight'].std(),
            'mean_Model_2_Weight': df_seg['Model_2_Weight'].mean(),
            'mean_Gate_Routing_Correct': df_seg['Gate_Routing_Correct'].mean()
        }

        for model_name, pred_col in [
            ('M1', 'Model_1_Pred_LSTM'),
            ('M2', 'Model_2_Pred_XGB'),
            ('Fused', 'Fused_Pred')
        ]:
            rmse, nse, mae, mape = compute_metrics_np(
                df_seg['True_Value'].values,
                df_seg[pred_col].values
            )

            # NSE is not meaningful for very short or constant segments.
            if not np.isfinite(nse):
                nse = np.nan

            row[f'RMSE_{model_name}'] = rmse
            row[f'NSE_{model_name}'] = nse
            row[f'MAE_{model_name}'] = mae
            row[f'MAPE_{model_name}'] = mape

        row['Fused_Better_Than_M1'] = row['RMSE_Fused'] < row['RMSE_M1']
        row['Fused_Better_Than_M2'] = row['RMSE_Fused'] < row['RMSE_M2']
        row['Fused_Best'] = (
                row['RMSE_Fused'] <= min(row['RMSE_M1'], row['RMSE_M2'])
        )

        segment_summary_rows.append(row)

    segment_summary_df = pd.DataFrame(segment_summary_rows)

    if not segment_summary_df.empty:
        segment_summary_df = segment_summary_df.sort_values(
            by=['segment_start', 'step']
        ).reset_index(drop=True)

    segment_summary_df.to_csv(
        os.path.join(fuse_output_dir_details, "Segment_Level_Test_Summary.csv"),
        index=False,
        encoding='utf-8-sig'
    )

    summary_rows = []

    def compute_step_metrics(df_sub, step_val, dataset_name):
        if len(df_sub) == 0:
            return
        rmse_lstm, nse_lstm, mae_lstm, mape_lstm = compute_metrics_np(df_sub['True_Value'].values, df_sub['Model_1_Pred_LSTM'].values)
        rmse_xgb, nse_xgb, mae_xgb, mape_xgb = compute_metrics_np(df_sub['True_Value'].values, df_sub['Model_2_Pred_XGB'].values)
        rmse_fused, nse_fused, mae_fused, mape_fused = compute_metrics_np(df_sub['True_Value'].values, df_sub['Fused_Pred'].values)
        for model, rmse, nse, mae, mape in [('LSTM', rmse_lstm, nse_lstm, mae_lstm, mape_lstm),
                                            ('XGB', rmse_xgb, nse_xgb, mae_xgb, mape_xgb),
                                            ('Fused', rmse_fused, nse_fused, mae_fused, mape_fused)]:
            summary_rows.append({
                'Step': step_val,
                'Dataset': dataset_name,
                'Model': model,
                'RMSE': rmse,
                'NSE': nse,
                'MAE': mae,
                'MAPE': mape
            })

    for ds_name, ds_mask in [('train', is_train), ('val', is_val_all), ('test', is_test)]:
        df_ds = out_df_all[ds_mask]
        compute_step_metrics(df_ds, 'All', ds_name)

    for step in sorted(out_df_all['step'].unique()):
        df_step = out_df_all[out_df_all['step'] == step]
        for ds_name, ds_mask_name in [('train', 'train'), ('val', 'val'), ('test', 'test')]:
            df_ds_step = df_step[df_step['Dataset'].str.lower() == ds_mask_name]
            compute_step_metrics(df_ds_step, step, ds_name)

    summary_df = pd.DataFrame(summary_rows)
    summary_df.to_csv(fuse_summary_filepath, index=False, encoding='utf-8-sig')

    test_all = summary_df[(summary_df['Step'] == 'All') & (summary_df['Dataset'] == 'test')]
    print("-" * 50)
    print(f"[Result] Gating fusion completed.")
    for _, row in test_all.iterrows():
        print(
            f"Test (All) - {row['Model']} -> RMSE: {row['RMSE']:.4f}, NSE: {row['NSE']:.4f}, MAE: {row['MAE']:.4f}, MAPE: {row['MAPE']:.4f}%")
    print("-" * 50)

    # Interpretation-only SHAP outputs are omitted in this release version.

    model_save_dir = get_model_save_dir(base_out_dir)
    gating_dir = os.path.join(model_save_dir, "gating_fusion")
    os.makedirs(gating_dir, exist_ok=True)
    torch.save(gating_net.state_dict(), os.path.join(gating_dir, "gating_net_best.pth"))
    with open(os.path.join(gating_dir, "scaler_gating.pkl"), 'wb') as f:
        pickle.dump(scaler_gate, f)



def write_model_manifest(target_well, all_wells, external_factors, base_out_dir):
    """Save a deployment manifest for impute_dataall.py.

    The manifest is intentionally explicit: prediction should never guess feature
    order, scaler paths, or available M1 steps from scattered output files.
    """
    model_save_dir = get_model_save_dir(base_out_dir)
    univ_dir = os.path.join(model_save_dir, "univariate")
    multiv_dir = os.path.join(model_save_dir, "multivariate")
    gating_dir = os.path.join(model_save_dir, "gating_fusion")

    m2_meta_path = os.path.join(multiv_dir, "m2_metadata.json")
    if os.path.exists(m2_meta_path):
        with open(m2_meta_path, "r", encoding="utf-8") as f:
            m2_meta = json.load(f)
    else:
        m2_meta = {
            "features_used": [],
            "feature_means": {},
            "model_path": os.path.abspath(os.path.join(multiv_dir, "model_Best_XGB_Comb.json")),
            "scaler_X_path": os.path.abspath(os.path.join(multiv_dir, "scaler_X_Best_XGB_Comb.pkl")),
            "scaler_y_path": os.path.abspath(os.path.join(multiv_dir, "scaler_y_Best_XGB_Comb.pkl"))
        }
        m2_file = os.path.join(base_out_dir, "LSTM_Results", "Best_Comb_M2.csv")
        if os.path.exists(m2_file):
            try:
                df2 = pd.read_csv(m2_file, nrows=1)
                m2_meta["features_used"] = str(df2["Features_Used"].iloc[0]).split("|")
            except Exception:
                pass

    # Record only M1 steps that were successfully saved.
    trained_steps = []
    m1_model_paths = {}
    for step in TS_FORECAST_STEPS:
        p = os.path.join(univ_dir, f"model_{target_well}_step_{step}.keras")
        if os.path.exists(p):
            trained_steps.append(int(step))
            m1_model_paths[str(step)] = os.path.abspath(p)

    date_col = "date"
    data_file_abs = os.path.abspath(GLOBAL_DATA_FILE)
    if os.path.exists(data_file_abs):
        try:
            cols = pd.read_excel(data_file_abs, nrows=0).columns.tolist()
            date_candidates = [c for c in cols if str(c).lower() == "date"]
            if date_candidates:
                date_col = date_candidates[0]
        except Exception:
            pass

    features_used = list(m2_meta.get("features_used", []))
    manifest = {
        "target_well": target_well,
        "date_col": date_col,
        "all_wells": list(all_wells),
        "external_factors": list(external_factors),
        "context_length": int(CONTEXT_LENGTH),
        "trained_steps": trained_steps,
        "global_config_steps": list(map(int, TS_FORECAST_STEPS)),
        "data_file": data_file_abs,
        "base_out_dir": os.path.abspath(base_out_dir),
        "m1_scaler_path": os.path.abspath(os.path.join(univ_dir, f"scaler_{target_well}.pkl")),
        "m1_model_dir": os.path.abspath(univ_dir),
        "m1_model_paths": m1_model_paths,
        "m2_model_path": os.path.abspath(m2_meta.get("model_path", os.path.join(multiv_dir, "model_Best_XGB_Comb.json"))),
        "m2_scaler_X_path": os.path.abspath(m2_meta.get("scaler_X_path", os.path.join(multiv_dir, "scaler_X_Best_XGB_Comb.pkl"))),
        "m2_scaler_y_path": os.path.abspath(m2_meta.get("scaler_y_path", os.path.join(multiv_dir, "scaler_y_Best_XGB_Comb.pkl"))),
        "m2_features_used": features_used,
        "m2_feature_means": m2_meta.get("feature_means", {}),
        "gating_model_path": os.path.abspath(os.path.join(gating_dir, "gating_net_best.pth")),
        "gating_scaler_path": os.path.abspath(os.path.join(gating_dir, "scaler_gating.pkl")),
        "gate_input_order": ["step", "abs_pos"] + features_used,
        "overlong_gap_rule": "If actual gap length exceeds the maximum trained M1 step, use M2 only for the entire gap."
    }

    manifest_path = os.path.join(base_out_dir, "model_manifest.json")
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
    print(f">>> Deployment manifest saved: {manifest_path}")

# =============================================================================
# Saved-model imputation code
# =============================================================================

import argparse
import os
import pickle
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import xgboost as xgb
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from tensorflow.keras.models import load_model



class FusionGatingNet(nn.Module):
    """Must match the architecture used in core_model.py."""

    def __init__(self, input_dim: int):
        super().__init__()
        self.net = nn.Sequential(
            nn.Linear(input_dim, 64),
            nn.ReLU(),
            nn.Dropout(0.2),
            nn.Linear(64, 32),
            nn.ReLU(),
            nn.Linear(32, 1),
            nn.Sigmoid(),
        )

    def forward(self, x):
        return self.net(x).squeeze(-1)


def resolve_path(path: str, manifest_dir: Path) -> str:
    """Resolve a model/scaler path saved in the manifest.

    Manifests may contain absolute paths from the training computer. If the project
    directory is moved, this function also tries paths relative to the current
    Output_<well> folder, especially the saved_models/... suffix.
    """
    if not path:
        return path

    p = Path(path)
    if p.exists():
        return str(p)

    parts = list(p.parts)
    if "saved_models" in parts:
        suffix = Path(*parts[parts.index("saved_models"):])
        candidate = manifest_dir / suffix
        if candidate.exists():
            return str(candidate)

    # Fallback: search by filename inside this well's output folder.
    matches = list(manifest_dir.rglob(p.name))
    if matches:
        return str(matches[0])

    candidate = Path.cwd() / path
    if candidate.exists():
        return str(candidate)

    return str(p)


class WellPredictor:
    def __init__(self, manifest_path: str):
        self.manifest_path = Path(manifest_path)
        self.manifest_dir = self.manifest_path.parent
        self.manifest = read_json(str(self.manifest_path))
        self.target_well = self.manifest["target_well"]
        self.context_length = int(self.manifest.get("context_length", 8))
        self.trained_steps = sorted([int(x) for x in self.manifest.get("trained_steps", [])])
        if not self.trained_steps:
            raise ValueError(f"{self.target_well}:  has no available trained_steps in the manifest.")
        self.max_step = max(self.trained_steps)

        self.m2_features_used: List[str] = list(self.manifest.get("m2_features_used", []))
        self.m2_feature_means: Dict[str, float] = {
            k: safe_float(v, 0.0) for k, v in self.manifest.get("m2_feature_means", {}).items()
        }

        self.device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

        self.m1_scaler = self._load_pickle("m1_scaler_path")
        self.m1_models: Dict[int, object] = {}

        self.m2_model = None
        self.m2_scaler_X = None
        self.m2_scaler_y = None
        self._load_m2()

        self.gate_model = None
        self.gate_scaler = None
        self._load_gate()

    def _path(self, key: str) -> str:
        return resolve_path(self.manifest.get(key, ""), self.manifest_dir)

    def _load_pickle(self, key: str):
        path = self._path(key)
        if not path or not os.path.exists(path):
            raise FileNotFoundError(f"{self.target_well}: Missing {key}: {path}")
        with open(path, "rb") as f:
            return pickle.load(f)

    def _load_m2(self) -> None:
        model_path = self._path("m2_model_path")
        scaler_x_path = self._path("m2_scaler_X_path")
        scaler_y_path = self._path("m2_scaler_y_path")
        if not self.m2_features_used:
            return
        if not (os.path.exists(model_path) and os.path.exists(scaler_x_path) and os.path.exists(scaler_y_path)):
            return
        model = xgb.XGBRegressor()
        model.load_model(model_path)
        self.m2_model = model
        with open(scaler_x_path, "rb") as f:
            self.m2_scaler_X = pickle.load(f)
        with open(scaler_y_path, "rb") as f:
            self.m2_scaler_y = pickle.load(f)

    def _load_gate(self) -> None:
        gate_model_path = self._path("gating_model_path")
        gate_scaler_path = self._path("gating_scaler_path")
        if not (os.path.exists(gate_model_path) and os.path.exists(gate_scaler_path)):
            return
        with open(gate_scaler_path, "rb") as f:
            self.gate_scaler = pickle.load(f)
        input_dim = len(self.manifest.get("gate_input_order", ["step", "abs_pos"] + self.m2_features_used))
        model = FusionGatingNet(input_dim=input_dim).to(self.device)
        state = torch.load(gate_model_path, map_location=self.device)
        model.load_state_dict(state)
        model.eval()
        self.gate_model = model

    @property
    def m2_available(self) -> bool:
        return self.m2_model is not None and self.m2_scaler_X is not None and self.m2_scaler_y is not None

    @property
    def gate_available(self) -> bool:
        return self.gate_model is not None and self.gate_scaler is not None

    def select_upper_step(self, length: int) -> Optional[int]:
        for step in self.trained_steps:
            if step >= length:
                return step
        return None

    def build_m2_raw_for_row(self, df_original: pd.DataFrame, row_idx: int) -> Tuple[List[float], int]:
        """Build raw M2/gate features in exactly the training feature order."""
        values: List[float] = []
        missing_feature_count = 0

        for feat in self.m2_features_used:
            if feat.endswith("_mask"):
                base = feat[: -len("_mask")]
                observed = base in df_original.columns and pd.notna(df_original.at[row_idx, base])
                values.append(1.0 if observed else 0.0)
            else:
                if feat in df_original.columns and pd.notna(df_original.at[row_idx, feat]):
                    values.append(safe_float(df_original.at[row_idx, feat], self.m2_feature_means.get(feat, 0.0)))
                else:
                    values.append(safe_float(self.m2_feature_means.get(feat, 0.0), 0.0))
                    missing_feature_count += 1

        return values, missing_feature_count

    def predict_m2_rows(self, df_original: pd.DataFrame, row_indices: List[int]) -> Tuple[np.ndarray, np.ndarray, List[int]]:
        if not self.m2_available:
            return np.full(len(row_indices), np.nan), np.empty((len(row_indices), 0)), [0] * len(row_indices)

        raw_rows = []
        missing_counts = []
        for r in row_indices:
            raw, cnt = self.build_m2_raw_for_row(df_original, r)
            raw_rows.append(raw)
            missing_counts.append(cnt)

        X_raw = np.asarray(raw_rows, dtype=float)
        X_scaled = self.m2_scaler_X.transform(X_raw)
        y_scaled = self.m2_model.predict(X_scaled).reshape(-1, 1)
        y = self.m2_scaler_y.inverse_transform(y_scaled).flatten()
        return y, X_raw, missing_counts

    def _load_m1_model(self, step: int):
        if step in self.m1_models:
            return self.m1_models[step]
        paths = self.manifest.get("m1_model_paths", {})
        model_path = paths.get(str(step))
        if model_path is None:
            model_path = os.path.join(self.manifest.get("m1_model_dir", ""), f"model_{self.target_well}_step_{step}.keras")
        model_path = resolve_path(model_path, self.manifest_dir)
        if not os.path.exists(model_path):
            raise FileNotFoundError(f"{self.target_well}: Missing M1 step={step} model: {model_path}")
        model = load_model(model_path)
        self.m1_models[step] = model
        return model

    def get_m1_context(
        self,
        df_original: pd.DataFrame,
        start_idx: int,
        actual_length: int,
        selected_step: int,
        overlong: bool,
    ) -> Tuple[Optional[np.ndarray], str, Optional[int]]:
        """Return M1 context only for gaps within the trained length range.

        The future context is aligned with the selected trained model length.
        Overlong gaps are handled entirely by M2.
        """
        if overlong or actual_length > self.max_step:
            return None, "M1_unavailable_gap_exceeds_max_step", None

        c = self.context_length
        past_start = start_idx - c
        if past_start < 0:
            return None, "M1_unavailable_not_enough_past_context", None

        future_start = start_idx + selected_step
        if future_start + c > len(df_original):
            return None, "M1_unavailable_not_enough_future_context", future_start

        target = self.target_well
        past = df_original.loc[past_start : start_idx - 1, target].to_numpy(dtype=float)
        future = df_original.loc[future_start : future_start + c - 1, target].to_numpy(dtype=float)

        if len(past) != c or len(future) != c:
            return None, "M1_unavailable_context_length_error", future_start
        if np.isnan(past).any():
            return None, "M1_unavailable_past_context_has_nan", future_start
        if np.isnan(future).any():
            return None, "M1_unavailable_future_context_has_nan", future_start

        return np.concatenate([past, future]), "M1_available", future_start

    def predict_m1_prefix(
        self,
        df_original: pd.DataFrame,
        start_idx: int,
        actual_length: int,
        selected_step: int,
        prefix_len: int,
        overlong: bool,
    ) -> Tuple[Optional[np.ndarray], str, Optional[int]]:
        context, reason, future_start = self.get_m1_context(
            df_original=df_original,
            start_idx=start_idx,
            actual_length=actual_length,
            selected_step=selected_step,
            overlong=overlong,
        )
        if context is None:
            return None, reason, future_start

        model = self._load_m1_model(selected_step)
        scaled = self.m1_scaler.transform(context.reshape(-1, 1)).reshape(1, self.context_length * 2, 1)
        pred_scaled = model.predict(scaled, verbose=0).reshape(-1, 1)
        pred = self.m1_scaler.inverse_transform(pred_scaled).flatten()
        return pred[:prefix_len], reason, future_start

    def predict_gate_weights(self, gate_step: int, x_m2_raw: np.ndarray, prefix_len: int) -> np.ndarray:
        if not self.gate_available:
            return np.full(prefix_len, np.nan)
        rows = []
        for k in range(prefix_len):
            abs_pos = k + 1
            rows.append([float(gate_step), float(abs_pos)] + list(x_m2_raw[k]))
        X_raw = np.asarray(rows, dtype=float)
        X_scaled = self.gate_scaler.transform(X_raw)
        with torch.no_grad():
            t = torch.tensor(X_scaled, dtype=torch.float32).to(self.device)
            w = self.gate_model(t).detach().cpu().numpy().flatten()
        return w


def load_predictors(manifest_root: str, wells: List[str]) -> Dict[str, WellPredictor]:
    predictors: Dict[str, WellPredictor] = {}
    root = Path(manifest_root)
    for well in wells:
        manifest_path = root / f"Output_{well}" / "model_manifest.json"
        if not manifest_path.exists():
            print(f"[Warning] Missing model manifest for {well}; skipping this well: {manifest_path}")
            continue
        predictors[well] = WellPredictor(str(manifest_path))
        print(f"[Info] Loaded {well}  model. Max M1 step={predictors[well].max_step}, M2={predictors[well].m2_available}, Gate={predictors[well].gate_available}")
    return predictors


def impute_one_well(
    df_original: pd.DataFrame,
    df_imputed: pd.DataFrame,
    date_col: str,
    predictor: WellPredictor,
) -> List[Dict]:
    target = predictor.target_well
    logs: List[Dict] = []
    segments = contiguous_nan_segments(df_original[target])

    if not segments:
        print(f"[Info] {target}: no missing segments.")
        return logs

    print(f"[Info] {target}: found {len(segments)} missing segments.")

    for seg_idx, seg in enumerate(segments, start=1):
        start = int(seg["start"])
        end = int(seg["end"])
        L = int(seg["length"])
        row_indices = list(range(start, end + 1))
        segment_id = f"{target}_seg_{seg_idx:04d}_{start}_{end}"

        m2_preds, x_m2_raw, missing_counts = predictor.predict_m2_rows(df_original, row_indices)
        m2_ok = predictor.m2_available and np.isfinite(m2_preds).any()

        overlong = L > predictor.max_step
        if overlong:
            selected_step = None
            fusion_len = 0
            gate_step = None
        else:
            selected_step = predictor.select_upper_step(L)
            fusion_len = L if selected_step is not None else 0
            gate_step = selected_step

        m1_preds = None
        m1_reason = "M1_unavailable_no_selected_step"
        future_context_start = None
        if selected_step is not None and fusion_len > 0:
            try:
                m1_preds, m1_reason, future_context_start = predictor.predict_m1_prefix(
                    df_original=df_original,
                    start_idx=start,
                    actual_length=L,
                    selected_step=selected_step,
                    prefix_len=fusion_len,
                    overlong=overlong,
                )
            except Exception as exc:
                m1_preds = None
                m1_reason = f"M1_error_{type(exc).__name__}: {exc}"

        m1_ok = m1_preds is not None and len(m1_preds) == fusion_len

        if m1_ok and m2_ok and predictor.gate_available:
            gate_weights = predictor.predict_gate_weights(gate_step=gate_step, x_m2_raw=x_m2_raw[:fusion_len], prefix_len=fusion_len)
        else:
            gate_weights = np.full(fusion_len, np.nan)

        for k, row_idx in enumerate(row_indices):
            abs_pos_real = k + 1
            in_fusion_prefix = k < fusion_len

            m1_pred = float(m1_preds[k]) if m1_ok and in_fusion_prefix else np.nan
            m2_pred = float(m2_preds[k]) if m2_ok and np.isfinite(m2_preds[k]) else np.nan
            w1 = float(gate_weights[k]) if in_fusion_prefix and np.isfinite(gate_weights[k]) else np.nan
            w2 = 1.0 - w1 if np.isfinite(w1) else np.nan

            if overlong:
                if np.isfinite(m2_pred):
                    final = m2_pred
                    method = "M2_only_overlong"
                    reason = "actual_length_exceeds_max_step_entire_gap_uses_M2"
                else:
                    final = np.nan
                    method = "Skipped"
                    reason = "overlong_gap_M2_unavailable"
            elif m1_ok and m2_ok and predictor.gate_available and np.isfinite(w1) and np.isfinite(m2_pred):
                final = w1 * m1_pred + (1.0 - w1) * m2_pred
                method = "Fused"
                reason = "M1_M2_gate_available"
            elif np.isfinite(m2_pred):
                final = m2_pred
                method = "M2_only"
                reason = m1_reason if not m1_ok else "Gate_unavailable_or_M1_M2_alignment_failed"
            elif np.isfinite(m1_pred):
                final = m1_pred
                method = "M1_only"
                reason = "M2_unavailable"
            else:
                final = np.nan
                method = "Skipped"
                reason = "M1_and_M2_unavailable"

            if np.isfinite(final):
                df_imputed.at[row_idx, target] = final

            logs.append(
                {
                    "target_well": target,
                    "row_index": row_idx,
                    "date": df_original.at[row_idx, date_col],
                    "segment_id": segment_id,
                    "segment_start_row": start,
                    "segment_end_row": end,
                    "segment_start_date": df_original.at[start, date_col],
                    "segment_end_date": df_original.at[end, date_col],
                    "missing_length_real": L,
                    "abs_pos_real": abs_pos_real,
                    "selected_m1_step": selected_step,
                    "gate_step_used": gate_step if in_fusion_prefix else np.nan,
                    "fusion_prefix_length": fusion_len,
                    "overlong_gap": bool(overlong),
                    "future_context_start_row_for_M1": future_context_start,
                    "m1_available": bool(m1_ok and in_fusion_prefix),
                    "m2_available": bool(np.isfinite(m2_pred)),
                    "gate_available": bool(predictor.gate_available and in_fusion_prefix),
                    "m1_pred": m1_pred,
                    "m2_pred": m2_pred,
                    "model_1_weight": w1,
                    "model_2_weight": w2,
                    "final_pred": final,
                    "method": method,
                    "reason": reason,
                    "features_used": "|".join(predictor.m2_features_used),
                    "missing_feature_count": int(missing_counts[k]) if k < len(missing_counts) else np.nan,
                }
            )

    return logs


def write_output_excel(
    output_path: str,
    df_imputed: pd.DataFrame,
    log_df: pd.DataFrame,
    original_columns: List[str],
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_imputed.to_excel(writer, sheet_name="imputed_data", index=False)
        log_df.to_excel(writer, sheet_name="imputation_log", index=False)

    wb = load_workbook(output_path)
    ws = wb["imputed_data"]
    ws.freeze_panes = "A2"

    fills = {
        "Fused": PatternFill("solid", fgColor="FFF2CC"),                 # light yellow
        "M2_only": PatternFill("solid", fgColor="D9EAF7"),               # light blue
        "M2_only_overlong": PatternFill("solid", fgColor="D9EAF7"),
        "M1_only": PatternFill("solid", fgColor="D9EAD3"),               # light green
        "Skipped": PatternFill("solid", fgColor="F4CCCC"),               # light red
    }

    col_to_excel_idx = {col: i + 1 for i, col in enumerate(original_columns)}
    if not log_df.empty:
        for _, row in log_df.iterrows():
            method = str(row.get("method", ""))
            if method == "Skipped":
                continue
            target = row.get("target_well")
            if target not in col_to_excel_idx:
                continue
            excel_row = int(row["row_index"]) + 2  # pandas row 0 -> Excel row 2 because header
            excel_col = col_to_excel_idx[target]
            ws.cell(row=excel_row, column=excel_col).fill = fills.get(method, fills.get("M2_only"))

    # Basic readability formatting
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for col_cells in ws.columns:
        col_letter = col_cells[0].column_letter
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:200])
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 22)

    if "imputation_log" in wb.sheetnames:
        log_ws = wb["imputation_log"]
        log_ws.freeze_panes = "A2"
        for cell in log_ws[1]:
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        for col_cells in log_ws.columns:
            col_letter = col_cells[0].column_letter
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:200])
            log_ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 35)

    wb.save(output_path)

# ==============================================================================
# Compact-release orchestration helpers
# ==============================================================================

from pathlib import Path as _Path
import subprocess as _subprocess
import time as _time
from typing import Sequence as _Sequence
from openpyxl import load_workbook as _load_workbook
from openpyxl.styles import Font as _Font, PatternFill as _PatternFill


def train_one_well(target_well: str, all_wells: list, external_factors: list, data_file: str, well_dir: str) -> None:
    """Train M1, M2, and the gating model for one target well and save a manifest."""
    global GLOBAL_DATA_FILE
    GLOBAL_DATA_FILE = os.path.abspath(data_file)

    if target_well not in all_wells:
        raise ValueError(f"Invalid target well: {target_well}. Available wells: {all_wells}")

    other_wells = [w for w in all_wells if w != target_well]
    dynamic_features = other_wells + list(external_factors)
    if not dynamic_features:
        raise ValueError(f"No dynamic predictors are available for {target_well}.")

    base_out_dir = os.path.abspath(well_dir)
    lstm_results_dir = os.path.join(base_out_dir, "LSTM_Results")
    fuse_output_dir_details = os.path.join(base_out_dir, "Fused_Details")
    os.makedirs(lstm_results_dir, exist_ok=True)
    os.makedirs(fuse_output_dir_details, exist_ok=True)
    fuse_summary_file = os.path.join(base_out_dir, "Fusion_Summary.csv")

    print("=" * 88)
    print(f"Training target well: {target_well}")
    print(f"Model directory: {base_out_dir}")
    print("=" * 88)

    step1_univariate_sequence(target_well, lstm_results_dir, base_out_dir)
    step2_multivariate_external(target_well, dynamic_features, lstm_results_dir, base_out_dir)
    step3_fusion(lstm_results_dir, base_out_dir, fuse_output_dir_details, fuse_summary_file)
    write_model_manifest(target_well, all_wells, external_factors, base_out_dir)


def train_all_wells(data_file: str, models_dir: str, stop_on_error: bool = False) -> None:
    """Train all well-specific models in independent subprocesses."""
    data_path = _Path(data_file).resolve()
    models_path = _Path(models_dir).resolve()
    models_path.mkdir(parents=True, exist_ok=True)

    if not data_path.exists():
        raise FileNotFoundError(f"Training dataset not found: {data_path}")

    header = clean_column_names(pd.read_excel(data_path, nrows=0))
    col_info = identify_columns(header)
    wells_str = ",".join(col_info.well_cols)
    external_str = ",".join(col_info.external_cols)

    for well in col_info.well_cols:
        print("\n" + "#" * 88)
        print(f"Launching single-well training process for: {well}")
        print("#" * 88)
        start = _time.time()
        cmd = [
            sys.executable,
            str(_Path(__file__).resolve()),
            "train-one",
            "--target", well,
            "--all-wells", wells_str,
            "--external", external_str,
            "--data", str(data_path),
            "--well-dir", str(models_path / well),
        ]
        try:
            _subprocess.run(cmd, check=True)
            print(f"[Done] {well} finished in {_time.time() - start:.2f} seconds.")
        except _subprocess.CalledProcessError as exc:
            print(f"[Error] Training failed for {well}, exit code {exc.returncode}.")
            if stop_on_error:
                raise


def _write_training_block_sheet(writer, summary: dict) -> None:
    rows = []
    for k, v in summary.items():
        if isinstance(v, list):
            v = ", ".join(map(str, v))
        rows.append({"item": k, "value": v})
    pd.DataFrame(rows).to_excel(writer, sheet_name="selected_training_block", index=False)


def collect_overall_performance(models_dir: str, training_block_summary_json: str, output_xlsx: str) -> None:
    """Create the compact performance summary: selected block + overall metrics only."""
    models_path = _Path(models_dir)
    rows = []
    model_name_map = {"LSTM": "M1_BiLSTM", "XGB": "M2_XGBoost", "Fused": "Fused"}

    for well_dir in sorted([p for p in models_path.iterdir() if p.is_dir()]):
        summary_csv = well_dir / "Fusion_Summary.csv"
        if not summary_csv.exists():
            print(f"[Warning] Missing Fusion_Summary.csv for {well_dir.name}")
            continue
        df = pd.read_csv(summary_csv)
        if "Step" not in df.columns:
            continue
        df = df[df["Step"].astype(str).str.lower() == "all"].copy()
        if df.empty:
            continue
        df.insert(0, "target_well", well_dir.name)
        df["Model"] = df["Model"].map(model_name_map).fillna(df["Model"])
        rows.append(df[["target_well", "Dataset", "Model", "RMSE", "NSE", "MAE", "MAPE"]])

    if rows:
        overall = pd.concat(rows, ignore_index=True)
        dataset_order = {"train": 0, "val": 1, "test": 2}
        model_order = {"M1_BiLSTM": 0, "M2_XGBoost": 1, "Fused": 2}
        overall["_d"] = overall["Dataset"].astype(str).str.lower().map(dataset_order).fillna(99)
        overall["_m"] = overall["Model"].map(model_order).fillna(99)
        overall = overall.sort_values(["target_well", "_d", "_m"]).drop(columns=["_d", "_m"])
        overall = overall.rename(columns={"Dataset": "dataset", "Model": "model"})
    else:
        overall = pd.DataFrame(columns=["target_well", "dataset", "model", "RMSE", "NSE", "MAE", "MAPE"])

    summary = read_json(training_block_summary_json) if os.path.exists(training_block_summary_json) else {}
    output_path = _Path(output_xlsx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_training_block_sheet(writer, summary)
        overall.to_excel(writer, sheet_name="overall_metrics", index=False)

    wb = _load_workbook(output_path)
    header_fill = _PatternFill("solid", fgColor="D9E1F2")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = _Font(bold=True)
        for col_cells in ws.columns:
            letter = col_cells[0].column_letter
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:200])
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 45)
    wb.save(output_path)


def cleanup_intermediate_files(models_dir: str) -> None:
    """Remove non-essential details while preserving deployable model files and manifests."""
    models_path = _Path(models_dir)
    for well_dir in models_path.iterdir() if models_path.exists() else []:
        if not well_dir.is_dir():
            continue
        for name in ["LSTM_Results", "Fused_Details"]:
            p = well_dir / name
            if p.exists():
                import shutil
                shutil.rmtree(p, ignore_errors=True)
        for name in ["Fusion_Summary.csv", "step_test_metrics.csv", "SHAP_Feature_Importance.csv"]:
            p = well_dir / name
            if p.exists():
                p.unlink()
    for cache in models_path.rglob("__pycache__") if models_path.exists() else []:
        import shutil
        shutil.rmtree(cache, ignore_errors=True)


# Override the earlier deployment loader with the compact release directory layout.
def load_predictors(models_dir: str, wells: list) -> dict:
    """Load one predictor per well from models_dir/<well>/model_manifest.json."""
    predictors = {}
    root = _Path(models_dir)
    missing = []
    for well in wells:
        candidates = [
            root / well / "model_manifest.json",
            root / f"Output_{well}" / "model_manifest.json",  # backward-compatible fallback
        ]
        manifest_path = next((p for p in candidates if p.exists()), None)
        if manifest_path is None:
            missing.append(well)
            continue
        predictors[well] = WellPredictor(str(manifest_path))
        print(
            f"[Info] Loaded {well}: max M1 step={predictors[well].max_step}, "
            f"M2={predictors[well].m2_available}, Gate={predictors[well].gate_available}"
        )
    if missing:
        raise FileNotFoundError(
            "Missing model manifests for wells: " + ", ".join(missing) +
            f". Expected files such as {root}/<well>/model_manifest.json. "
            "Run groundwater_imputation_first_run.py first or check --models-dir."
        )
    return predictors


def run_saved_model_imputation(input_file: str, output_file: str, models_dir: str, external_cols: list | None = None) -> pd.DataFrame:
    """Impute an input Excel file using already saved models."""
    input_path = _Path(input_file).resolve()
    output_path = _Path(output_file).resolve()
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    df_raw = clean_column_names(pd.read_excel(input_path))
    col_info = identify_columns(df_raw, external_cols=external_cols)
    df_original = ensure_datetime_sorted(df_raw, col_info.date_col)
    df_imputed = df_original.copy()

    wells_with_missing = [w for w in col_info.well_cols if df_original[w].isna().any()]
    predictors = load_predictors(models_dir, col_info.well_cols)
    all_logs = []
    for well in col_info.well_cols:
        all_logs.extend(impute_one_well(df_original, df_imputed, col_info.date_col, predictors[well]))

    log_df = pd.DataFrame(all_logs)
    if wells_with_missing and log_df.empty:
        raise RuntimeError(
            "Missing well values were detected, but no imputation log was produced. "
            "Please check model files, column names, and missing-value encoding."
        )

    write_output_excel(str(output_path), df_imputed, log_df, list(df_imputed.columns))
    print(f"[Done] Imputed dataset written to: {output_path}")
    print(f"[Done] Imputation log rows: {len(log_df)}")
    return log_df


def _main_cli() -> None:
    parser = argparse.ArgumentParser(description="Internal training/imputation commands for the compact groundwater release.")
    sub = parser.add_subparsers(dest="command", required=True)

    p_train = sub.add_parser("train-one", help="Train one target well.")
    p_train.add_argument("--target", required=True)
    p_train.add_argument("--all-wells", required=True)
    p_train.add_argument("--external", default="")
    p_train.add_argument("--data", required=True)
    p_train.add_argument("--well-dir", required=True)

    p_imp = sub.add_parser("impute", help="Impute using saved models.")
    p_imp.add_argument("--input", required=True)
    p_imp.add_argument("--output", required=True)
    p_imp.add_argument("--models-dir", required=True)
    p_imp.add_argument("--external-cols", default=None)

    args = parser.parse_args()
    if args.command == "train-one":
        all_wells = [x for x in args.all_wells.split(",") if x]
        external = [x for x in args.external.split(",") if x]
        train_one_well(args.target, all_wells, external, args.data, args.well_dir)
    elif args.command == "impute":
        ext = [x.strip() for x in args.external_cols.split(",") if x.strip()] if args.external_cols else None
        run_saved_model_imputation(args.input, args.output, args.models_dir, external_cols=ext)


if __name__ == "__main__":
    _main_cli()
